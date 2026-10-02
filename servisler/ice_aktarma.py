"""CSV/Excel içe aktarma çekirdeği (Adım 4b-1, `docs/adim-4b-tasarim.md`, K29–K39). Saf: veritabanı ve FastAPI yok.

Akış: dosya_oku → esleme_tahmin_et (ya da kullanıcının eşleştirmesi) → esleme_dogrula → satirlari_hazirla.
Beyaz liste (K34): yalnızca eşleştirilen alanlar kayıtlara kopyalanır; diğer sütunlar bu modülden çıkmaz. Hassas
başlıklı sütunların hücreleri okunurken atılır. Hata ve uyarı metinleri hiçbir hücre DEĞERİ içermez; yalnızca satır
no, alan ve sebep (ad, telefon, e-posta değerleri hiçbir çıktıya düşmez).
Para yalnızca Decimal; xlsx'in ondalıklı sayı hücreleri okunurken Decimal(repr(x)) ile çevrilir, float dışarı çıkmaz.
"""

import csv
import hashlib
import io
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field, replace
from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

from servisler.telefon import telefon_normallestir

AZAMI_BAYT = 5 * 1024 * 1024
AZAMI_SATIR = 50_000
TURLER = ("uyeler", "paketler", "girisler")

# K31: normalleştirilmiş eş anlamlılar. "id" yalnızca başlığın tamamıysa ve dosyada başka kimlik sütunu yoksa.
_UYE_KIMLIK = ("uye no", "uye id", "uye kodu", "uye numarasi", "musteri no", "musteri id", "musteri kodu", "member id",
               "id")
ESANLAMLILAR: dict[str, dict[str, tuple[str, ...]]] = {
    "uyeler": {
        "uye_kimlik": _UYE_KIMLIK,
        "ad_soyad": ("ad soyad", "adi soyadi", "uye adi", "uye adi soyadi", "musteri adi", "isim", "isim soyisim",
                     "name", "full name"),
        "ad": ("ad", "adi", "first name"),
        "soyad": ("soyad", "soyadi", "last name"),
        "telefon": ("telefon", "tel", "gsm", "cep", "cep telefonu", "telefon no", "phone", "mobile"),
        "eposta": ("e posta", "eposta", "email", "e mail", "mail"),
    },
    "paketler": {
        "uye_kimlik": _UYE_KIMLIK,
        "paket_kimlik": ("paket no", "paket id", "uyelik no", "satis no"),
        "paket_adi": ("paket", "paket adi", "uyelik", "uyelik tipi", "uyelik turu", "package", "plan"),
        "baslangic": ("baslangic", "baslangic tarihi", "satis tarihi", "start", "start date"),
        "bitis": ("bitis", "bitis tarihi", "son kullanma", "son gecerlilik", "end", "end date"),
        "giris_hakki": ("giris hakki", "seans", "seans sayisi", "ders hakki", "ders sayisi", "kredi", "sessions"),
        "ucret": ("ucret", "tutar", "fiyat", "odenen", "odenen tutar", "price", "amount",
                  "fiyati", "paket fiyati", "uyelik ucreti", "ucreti", "tutari"),
        "durum": ("durum", "status", "durumu", "uyelik durumu", "paket durumu"),
    },
    "girisler": {
        "uye_kimlik": _UYE_KIMLIK,
        "tarih": ("tarih", "giris tarihi", "ziyaret tarihi", "ders tarihi", "check in", "checkin", "date"),
        "saat": ("saat", "giris saati", "time"),
        "giris_kimlik": ("giris no", "giris id", "kayit no"),
        "tutar": ("tutar", "ucret", "amount"),
        "durum": ("durum", "status"),
    },
}
ZORUNLU = {
    "uyeler": ("uye_kimlik", "ad_soyad"),
    "paketler": ("uye_kimlik", "paket_adi", "baslangic", "ucret"),
    "girisler": ("uye_kimlik", "tarih"),
}
_KIMLIK_ALANLARI = ("uye_kimlik", "paket_kimlik", "giris_kimlik")
_KISISEL_ALANLAR = ("ad_soyad", "ad", "soyad", "telefon", "eposta")     # --anonim'de hiç okunmaz (K35)
# "Kalan Seans" gibi başlıklar toplam hak değildir: bu kelimelerden biri geçen başlık giris_hakki'na eşleşmez.
_GIRIS_HAKKI_DISI = ("kalan", "kullanilan", "kullanilmis", "harcanan")
HASSAS_IFADELER = ("tc", "kimlik no", "kan", "biyometri", "parmak", "saglik", "hastalik", "ilac")

PAKET_DURUMLARI = {
    **dict.fromkeys(("aktif", "devam", "devam ediyor", "active"), "aktif"),
    **dict.fromkeys(("bitti", "sona erdi", "suresi doldu", "expired", "tamamlandi"), "bitti"),
    **dict.fromkeys(("iptal", "iptal edildi", "cancelled"), "iptal"),
    **dict.fromkeys(("donduruldu", "dondu", "frozen"), "donduruldu"),
}
GIRIS_DURUMLARI = {
    **dict.fromkeys(("geldi", "tamamlandi", "katildi"), "tamamlandi"),
    "iptal": "iptal",
    **dict.fromkeys(("gelmedi", "no show"), "gelmedi"),
}

UYARI_ACIKLAMALARI = {
    "binlik_nokta": "'1.250' biçimindeki tutarlar binlik ayırıcılı okundu (1.250 = 1250)",
    "telefon_gecersiz": "geçersiz telefon boş bırakıldı",
    "telefon_tekrar": "aynı telefon dosyada birden çok üyede; ilki dışındakiler boş bırakıldı",
    "telefon_cakisma": "telefon veritabanında başka bir üyede; boş bırakıldı",
    "eposta_gecersiz": "'@' içermeyen e-posta boş bırakıldı",
    "tekrar_satir": "aynı türetilmiş kimlikli tekrar satır tek kayıt sayıldı",
}
CIFT_SAYIM_UYARISI = "Paket geliri ve giriş tutarı birlikte: ciro iki kez sayılabilir"

_EXCEL_SIFIR = date(1899, 12, 30)
_VARSAYILAN_SAAT = time(12)
_GELECEK_PAYI = timedelta(minutes=5)
_KAYNAK_DESENI = re.compile(r"^[a-z0-9_-]{1,40}$")


class IceAktarmaHatasi(Exception):
    pass


class DosyaHatasi(IceAktarmaHatasi):
    """Dosya okunamıyor veya sınırları aşıyor; hiçbir satır işlenmez."""


class EslemeHatasi(IceAktarmaHatasi):
    def __init__(self, sorunlar: list[str]):
        super().__init__("; ".join(sorunlar))
        self.sorunlar = sorunlar


class DegerHatasi(IceAktarmaHatasi):
    """Tek bir hücre çözülemedi. Mesaj yalnızca sebeptir; hücre değeri asla eklenmez."""


# ---------------------------------------------------------------------------------------------
# Okuma (K30, K39)
# ---------------------------------------------------------------------------------------------
@dataclass(frozen=True)
class Tablo:
    basliklar: list[str]
    satirlar: list[tuple[int, tuple]]           # (dosyadaki satır no, hücreler); hücreler başlık sayısı kadar
    bicim: str                                  # 'csv' | 'xlsx'
    kodlama: str | None = None
    ayirici: str | None = None


def dosya_oku(icerik: bytes, dosya_adi: str) -> Tablo:
    if len(icerik) > AZAMI_BAYT:
        raise DosyaHatasi(f"Dosya {AZAMI_BAYT // (1024 * 1024)} MB sınırını aşıyor")
    uzanti = Path(dosya_adi).suffix.lower()
    if uzanti == ".xlsx":
        return _xlsx_oku(icerik)
    if uzanti in (".csv", ".txt", ".tsv"):
        return _csv_oku(icerik)
    raise DosyaHatasi("Desteklenmeyen dosya türü; .csv veya .xlsx verin")


def _kodlama_coz(icerik: bytes) -> tuple[str, str]:
    for kodlama in ("utf-8-sig", "cp1254"):
        try:
            return icerik.decode(kodlama), kodlama
        except UnicodeDecodeError:
            continue
    raise DosyaHatasi("Dosya kodlaması çözülemedi (utf-8 ve cp1254 denendi)")


def _ayirici_bul(metin: str) -> str:
    ornek = [s for s in metin.splitlines() if s.strip()][:5]
    try:
        return csv.Sniffer().sniff("\n".join(ornek), delimiters=";,\t").delimiter
    except csv.Error:                           # tek sütun ya da tutarsız: başlık satırında en sık aday
        ilk = ornek[0] if ornek else ""
        return max(";,\t", key=ilk.count)


def _csv_oku(icerik: bytes) -> Tablo:
    metin, kodlama = _kodlama_coz(icerik)
    ayirici = _ayirici_bul(metin)
    okuyucu = csv.reader(io.StringIO(metin, newline=""), delimiter=ayirici)
    try:
        basliklar, satirlar = _tabloya_cevir(((okuyucu.line_num, satir) for satir in okuyucu))
    except csv.Error:
        raise DosyaHatasi(f"CSV okunamadı (satır {okuyucu.line_num})") from None
    return Tablo(basliklar, satirlar, "csv", kodlama, ayirici)


def _xlsx_oku(icerik: bytes) -> Tablo:
    from openpyxl import load_workbook          # yalnızca burada: uygulama açılışı openpyxl'e bağlı değil

    try:
        kitap = load_workbook(io.BytesIO(icerik), read_only=True, data_only=True)
    except Exception:
        raise DosyaHatasi("xlsx dosyası okunamadı") from None
    try:
        sayfa = kitap.worksheets[0]
        basliklar, satirlar = _tabloya_cevir(enumerate(sayfa.iter_rows(values_only=True), start=1))
    finally:
        kitap.close()
    return Tablo(basliklar, satirlar, "xlsx")


def _hucre(deger):
    if isinstance(deger, str):
        deger = deger.strip()
        return deger or None
    if isinstance(deger, bool):
        return str(deger)
    if isinstance(deger, float):
        return Decimal(repr(deger))
    return deger


def _tabloya_cevir(kaynak) -> tuple[list[str], list[tuple[int, tuple]]]:
    """Boş satırları atlar; ilk dolu satır başlıktır. Hassas başlıklı sütunların hücreleri hiç saklanmaz (K34)."""
    basliklar: list[str] | None = None
    hassas: set[int] = set()
    satirlar: list[tuple[int, tuple]] = []
    for satir_no, ham in kaynak:
        hucreler = [_hucre(h) for h in ham]
        if all(h is None for h in hucreler):
            continue
        if basliklar is None:
            basliklar = ["" if h is None else str(h) for h in hucreler]
            hassas = {i for i, b in enumerate(basliklar) if hassas_mi(b)}
            continue
        if len(satirlar) >= AZAMI_SATIR:
            raise DosyaHatasi(f"Dosya {AZAMI_SATIR} veri satırı sınırını aşıyor")
        hucreler = (hucreler + [None] * len(basliklar))[:len(basliklar)]
        satirlar.append((satir_no, tuple(None if i in hassas else h for i, h in enumerate(hucreler))))
    if basliklar is None:
        raise DosyaHatasi("Dosya boş")
    return basliklar, satirlar


# ---------------------------------------------------------------------------------------------
# Eşleştirme (K31, K34)
# ---------------------------------------------------------------------------------------------
_TURKCE_ASCII = str.maketrans("şğıöüç", "sgiouc")
_AYIRICI_KARAKTER = re.compile(r"[^0-9a-z]+")


def baslik_normallestir(baslik: str) -> str:
    kucuk = baslik.replace("İ", "i").replace("I", "ı").lower().translate(_TURKCE_ASCII)
    return _AYIRICI_KARAKTER.sub(" ", kucuk).strip()


def hassas_mi(baslik: str) -> bool:
    norm = baslik_normallestir(baslik)
    return any(ifade in norm for ifade in HASSAS_IFADELER) or " t c " in f" {norm} "


@dataclass
class Esleme:
    alanlar: dict[str, str]                     # alan → dosyadaki özgün başlık
    belirsizler: list[str] = field(default_factory=list)
    eslesmeyenler: list[str] = field(default_factory=list)
    hassaslar: list[str] = field(default_factory=list)


def _geciyor(es: str, norm: str) -> bool:
    """Eş anlamlı başlıkta tam kelime dizisi olarak geçiyor mu. 4 harften kısa olanlar ("ad", "adi", "tel", "cep")
    yalnızca başlığın başında: "Cep Tel" telefondur ama "Paket Adı" üyenin adı değildir."""
    if len(es) < 4:
        return f"{norm} ".startswith(f"{es} ")
    return f" {es} " in f" {norm} "


def _en_iyi_alanlar(tur: str, norm: str) -> tuple[list[str], int]:
    """Başlıkta geçen en uzun eş anlamlının alan(lar)ı ve uzunluğu."""
    alanlar: list[str] = []
    en_uzun = 0
    kelimeler = norm.split()
    for alan, liste in ESANLAMLILAR[tur].items():
        if alan == "giris_hakki" and any(k in kelimeler for k in _GIRIS_HAKKI_DISI):
            continue
        for es in liste:
            if es == "id" or not _geciyor(es, norm):
                continue
            if len(es) > en_uzun:
                alanlar, en_uzun = [alan], len(es)
            elif len(es) == en_uzun and alan not in alanlar:
                alanlar.append(alan)
    return alanlar, en_uzun


def esleme_tahmin_et(tur: str, basliklar: list[str]) -> Esleme:
    sonuc = Esleme({})
    adaylar: dict[str, list[tuple[tuple[bool, int], str]]] = defaultdict(list)   # alan → [((tam mı, uzunluk), başlık)]
    id_basliklari: list[str] = []
    kimlik_sutunu_var = False
    for baslik in basliklar:
        norm = baslik_normallestir(baslik)
        if not norm:
            continue
        if hassas_mi(baslik):
            sonuc.hassaslar.append(baslik)
            continue
        if norm == "id":
            id_basliklari.append(baslik)
            continue
        alanlar, uzunluk = _en_iyi_alanlar(tur, norm)
        kimlik_sutunu_var |= any(a in _KIMLIK_ALANLARI for a in alanlar)
        if not alanlar:
            sonuc.eslesmeyenler.append(baslik)
        elif len(alanlar) > 1:
            sonuc.belirsizler.append(baslik)
        else:                                    # aynı alana birden çok başlık: tam eşleşme, sonra uzunluk
            adaylar[alanlar[0]].append(((len(norm) == uzunluk, uzunluk), baslik))
    for alan, liste in adaylar.items():
        en_iyi = max(o for o, _ in liste)
        kazananlar = [b for o, b in liste if o == en_iyi]
        if len(kazananlar) == 1:
            sonuc.alanlar[alan] = kazananlar[0]
        else:
            sonuc.belirsizler.extend(kazananlar)
        sonuc.eslesmeyenler.extend(b for o, b in liste if o < en_iyi)
    if len(id_basliklari) == 1 and not kimlik_sutunu_var:
        sonuc.alanlar["uye_kimlik"] = id_basliklari[0]
    else:
        sonuc.eslesmeyenler.extend(id_basliklari)
    sonuc.alanlar = {a: sonuc.alanlar[a] for a in ESANLAMLILAR[tur] if a in sonuc.alanlar}   # sözlük sırası
    return sonuc


def esleme_dogrula(tur: str, alanlar: dict[str, str | None], basliklar: list[str],
                   anonim: bool = False) -> dict[str, str]:
    """Boş (None) alanları atar; kişisel alanlar --anonim'de düşer. Sorun varsa EslemeHatasi (tüm sorunlarla)."""
    if tur not in ESANLAMLILAR:
        raise EslemeHatasi([f"bilinmeyen dosya türü: {tur}"])
    sorunlar: list[str] = []
    temiz: dict[str, str] = {}
    for alan, baslik in alanlar.items():
        if baslik is None or baslik == "":
            continue
        if alan not in ESANLAMLILAR[tur]:
            sorunlar.append(f"bilinmeyen alan: {alan}")
        elif not isinstance(baslik, str) or baslik not in basliklar:
            sorunlar.append(f"{alan}: dosyada olmayan başlık")
        elif hassas_mi(baslik):
            sorunlar.append(f"{alan}: hassas alan, içe aktarılmaz ({baslik!r})")
        elif not (anonim and alan in _KISISEL_ALANLAR):
            temiz[alan] = baslik
    for alan in ZORUNLU[tur]:
        if alan == "ad_soyad" and (anonim or "ad" in temiz):
            continue
        if alan not in temiz:
            sorunlar.append(f"zorunlu alan eşleşmedi: {alan}")
    if sorunlar:
        raise EslemeHatasi(sorunlar)
    return temiz


# ---------------------------------------------------------------------------------------------
# Dönüştürücüler (K37)
# ---------------------------------------------------------------------------------------------
_GUN_AY_YIL = re.compile(r"^(\d{1,2})([./-])(\d{1,2})\2(\d{4}|\d{2})(?:[ T]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$")
_YIL_AY_GUN = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$")
_SAAT = re.compile(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$")


def _saat_kur(saat, dakika, saniye) -> time:
    try:
        return time(int(saat), int(dakika), int(saniye or 0))
    except ValueError:
        raise DegerHatasi("geçersiz saat") from None


def _gun_kesri_saat(kesir: Decimal) -> time | None:
    saniye = min(int((kesir * 86_400).to_integral_value()), 86_399)
    return time(saniye // 3600, saniye % 3600 // 60, saniye % 60) if saniye else None


def _seri_tarih(n: Decimal) -> tuple[date, time | None]:
    """Excel seri tarihi: 1899-12-30 + n gün; kesirli kısım günün saati."""
    if not 20_000 <= n <= 80_000:
        raise DegerHatasi("tanınmayan tarih biçimi")
    gun = int(n)
    return _EXCEL_SIFIR + timedelta(days=gun), _gun_kesri_saat(n - gun)


def tarih_coz(deger) -> tuple[date, time | None]:
    """(gün, saat ya da None). Biçimler: gg.aa.yyyy, gg/aa/yyyy, gg-aa-yyyy, yyyy-aa-gg (+ SS:DD[:ss]), xlsx tarih
    hücresi, Excel seri sayısı. Ay > 12 → hata (aa/gg desteklenmez). xlsx'te 00:00 saati 'saatsiz' sayılır."""
    if deger is None:
        raise DegerHatasi("boş")
    if isinstance(deger, datetime):
        return deger.date(), (None if deger.time() == time(0) else deger.time().replace(microsecond=0))
    if isinstance(deger, date):
        return deger, None
    if isinstance(deger, (int, Decimal)):
        return _seri_tarih(Decimal(deger))
    metin = str(deger).strip()
    if re.fullmatch(r"\d{5}", metin):
        return _seri_tarih(Decimal(metin))
    if m := _GUN_AY_YIL.match(metin):
        gun, _, ay, yil, sa, dk, sn = m.groups()
    elif m := _YIL_AY_GUN.match(metin):
        yil, ay, gun, sa, dk, sn = m.groups()
    else:
        raise DegerHatasi("tanınmayan tarih biçimi")
    if int(ay) > 12 or int(ay) == 0:
        raise DegerHatasi("geçersiz ay (gün/ay sırası: gg.aa.yyyy beklenir)")
    yil = 2000 + int(yil) if len(yil) == 2 else int(yil)
    try:
        gun_tarihi = date(yil, int(ay), int(gun))
    except ValueError:
        raise DegerHatasi("geçersiz tarih") from None
    return gun_tarihi, (_saat_kur(sa, dk, sn) if sa is not None else None)


def saat_coz(deger) -> time | None:
    if deger is None:
        return None
    if isinstance(deger, datetime):
        return deger.time().replace(microsecond=0)
    if isinstance(deger, time):
        return deger.replace(microsecond=0)
    if isinstance(deger, (int, Decimal)) and 0 <= deger < 1:          # Excel saat kesri
        return _gun_kesri_saat(Decimal(deger)) or time(0)
    if m := _SAAT.match(str(deger).strip()):
        return _saat_kur(*m.groups())
    raise DegerHatasi("tanınmayan saat biçimi")


def _yerel_zaman(gun_saat: tuple[date, time | None], saat: time | None, dilim: ZoneInfo,
                 simdi: datetime) -> datetime:
    gun, tarihteki_saat = gun_saat
    zaman = datetime.combine(gun, saat or tarihteki_saat or _VARSAYILAN_SAAT, tzinfo=dilim).astimezone(UTC)
    if zaman > simdi + _GELECEK_PAYI:
        raise DegerHatasi("gelecek tarihli giriş")
    return zaman


def zaman_coz(tarih_degeri, saat_degeri, dilim: ZoneInfo, simdi: datetime | None = None) -> datetime:
    """Girişin UTC zamanı. Saat sütunu doluysa o, değilse tarihteki saat, o da yoksa yerel 12:00.
    Şu andan 5 dakikadan ileriyse hata."""
    return _yerel_zaman(tarih_coz(tarih_degeri), saat_coz(saat_degeri), dilim, simdi or datetime.now(UTC))


_BINLIK_NOKTA = re.compile(r"^\d{1,3}(\.\d{3})+$")
_SAYI = re.compile(r"^\d+(\.\d+)?$")


def para_coz(deger, uyarilar: Counter | None = None) -> Decimal:
    """Tutarı Decimal olarak döndürür (yuvarlamaz). 2'den fazla ondalık ya da negatif → hata. float kullanılmaz."""
    if deger is None:
        raise DegerHatasi("boş")
    if isinstance(deger, bool):
        raise DegerHatasi("tanınmayan tutar")
    if isinstance(deger, (int, Decimal)):
        tutar = Decimal(deger)
    else:
        metin = re.sub(r"\s|₺|(?i:tl)", "", str(deger))
        if metin.startswith("-"):
            raise DegerHatasi("negatif tutar")
        if "." in metin and "," in metin:
            ondalik = "," if metin.rfind(",") > metin.rfind(".") else "."
            metin = metin.replace("." if ondalik == "," else ",", "").replace(",", ".")
        elif "," in metin:
            metin = metin.replace(",", ".")
        elif _BINLIK_NOKTA.match(metin):
            metin = metin.replace(".", "")
            if uyarilar is not None:
                uyarilar["binlik_nokta"] += 1
        if not _SAYI.match(metin):
            raise DegerHatasi("tanınmayan tutar")
        try:
            tutar = Decimal(metin)
        except InvalidOperation:
            raise DegerHatasi("tanınmayan tutar") from None
    if not tutar.is_finite():
        raise DegerHatasi("tanınmayan tutar")
    if tutar < 0:
        raise DegerHatasi("negatif tutar")
    if -tutar.as_tuple().exponent > 2:
        raise DegerHatasi("2'den fazla ondalık hane")
    return tutar


def durum_coz(tur: str, deger) -> str | None:
    """Paket ya da giriş durumunu veritabanı değerine çevirir. Boş değer: paket → None (çıkarım), giriş →
    'tamamlandi'. Tanınmayan değer → hata."""
    norm = baslik_normallestir(str(deger)) if deger is not None else ""
    if not norm:
        return None if tur == "paketler" else "tamamlandi"
    sozluk = PAKET_DURUMLARI if tur == "paketler" else GIRIS_DURUMLARI
    if norm not in sozluk:
        raise DegerHatasi("tanınmayan durum")
    return sozluk[norm]


def _metin(deger) -> str | None:
    """Kimlik/telefon gibi metinler: xlsx'in tam sayı hücresi '1001.0' değil '1001' olur."""
    if deger is None:
        return None
    if isinstance(deger, Decimal) and deger == deger.to_integral_value():
        return str(int(deger))
    return str(deger).strip() or None


def _tam_sayi(deger) -> int:
    if isinstance(deger, int) and not isinstance(deger, bool):
        return deger
    if isinstance(deger, Decimal) and deger == deger.to_integral_value():
        return int(deger)
    if isinstance(deger, str) and re.fullmatch(r"\d+", deger.strip()):
        return int(deger)
    raise DegerHatasi("tam sayı değil")


# ---------------------------------------------------------------------------------------------
# Satırların hazırlanması (K33–K38)
# ---------------------------------------------------------------------------------------------
@dataclass(frozen=True)
class SatirHatasi:
    satir: int
    alan: str | None
    sebep: str

    def __str__(self) -> str:
        return f"satır {self.satir}" + (f", {self.alan}" if self.alan else "") + f": {self.sebep}"


@dataclass(frozen=True)
class UyeKaydi:
    satir: int
    dis_kimlik: str
    ad_soyad: str
    telefon_e164: str | None
    eposta: str | None


@dataclass(frozen=True)
class PaketKaydi:
    satir: int
    dis_kimlik: str
    uye_kimlik: str
    tur: str                                    # 'sure' | 'giris'
    ad: str
    baslangic_tarihi: date
    bitis_tarihi: date | None
    giris_hakki: int | None
    ucret: Decimal
    durum: str
    onceki_dis_kimlik: str | None               # yenileme zinciri: aynı üyenin bir önceki paketi


@dataclass(frozen=True)
class GirisKaydi:
    satir: int
    dis_kimlik: str
    uye_kimlik: str
    ziyaret_zamani: datetime                    # UTC
    durum: str
    toplam_tutar: Decimal


@dataclass
class Hazirlik:
    tur: str
    dis_kaynak: str
    okunan: int
    alanlar: tuple[str, ...] = ()               # okunan (eşleşmiş, doğrulanmış) alanlar
    anonim: bool = False
    kayitlar: list = field(default_factory=list)
    hatalar: list[SatirHatasi] = field(default_factory=list)
    uyarilar: Counter = field(default_factory=Counter)

    @property
    def hatali(self) -> int:
        return len({h.satir for h in self.hatalar})

    @property
    def gecerli(self) -> int:
        return len(self.kayitlar)

    @property
    def esik_asildi(self) -> bool:
        """K38: hatalı satır oranı %10'un üzerinde."""
        return self.hatali * 10 > self.okunan


def dis_kaynak_olustur(kaynak: str) -> str:
    if not isinstance(kaynak, str) or not _KAYNAK_DESENI.match(kaynak):
        raise ValueError("Kaynak adı yalnızca a-z, 0-9, '_' ve '-' içerebilir (1–40 karakter)")
    return f"csv:{kaynak}"


def _ozet(*parcalar: str) -> str:
    return hashlib.sha256("|".join(parcalar).encode("utf-8")).hexdigest()


def anonim_ad(dis_kaynak: str, uye_kimlik: str) -> str:
    return "Üye " + _ozet(dis_kaynak, uye_kimlik)[:6]


def ciro_cift_sayim_uyarisi(giris_alanlari: dict[str, str] | None, paket_dosyasi_var: bool) -> str | None:
    """K36: giriş tutarı okunuyorken aynı çalıştırmada paket dosyası da varsa uyarı metni."""
    return CIFT_SAYIM_UYARISI if paket_dosyasi_var and giris_alanlari and "tutar" in giris_alanlari else None


def satirlari_hazirla(tur: str, tablo: Tablo, esleme: dict[str, str | None], saat_dilimi: str, bugun: date,
                      anonim: bool, dis_kaynak: str, *, simdi: datetime | None = None,
                      bilinen_uyeler: set[str] | None = None) -> Hazirlik:
    """Geçerli kayıtlar + satır hataları + uyarı sayaçları. Yalnızca eşleşen alanlar okunur (K34).

    bilinen_uyeler verilirse paket/girişteki üye kimliği bu kümede olmalıdır (bu çalıştırmanın üyeleri ∪ aynı
    dis_kaynak'lı veritabanı üyeleri; K38); verilmezse denetlenmez.
    """
    alanlar = esleme_dogrula(tur, esleme, tablo.basliklar, anonim)
    sutun = {alan: tablo.basliklar.index(baslik) for alan, baslik in alanlar.items()}
    h = Hazirlik(tur=tur, dis_kaynak=dis_kaynak, okunan=len(tablo.satirlar), alanlar=tuple(alanlar), anonim=anonim)
    satir_hazirla = {"uyeler": _uye, "paketler": _paket, "girisler": _giris}[tur]
    baglam = _Baglam(ZoneInfo(saat_dilimi), bugun, simdi or datetime.now(UTC), anonim, dis_kaynak, bilinen_uyeler)
    gorulen: dict[str, int] = {}                     # dis_kimlik → ilk satır no
    for satir_no, hucreler in tablo.satirlar:
        deger = {alan: hucreler[i] for alan, i in sutun.items()}      # beyaz liste: yalnızca eşleşen alanlar
        hatalar: list[SatirHatasi] = []
        kayit = satir_hazirla(satir_no, deger, baglam, hatalar, h.uyarilar)
        if hatalar or kayit is None:
            h.hatalar.extend(hatalar)
            continue
        if kayit.dis_kimlik in gorulen:
            if tur != "uyeler" and _metin(deger.get(_kimlik_alani(tur))) is None:     # türetilmiş kimlik
                h.uyarilar["tekrar_satir"] += 1
            else:
                h.hatalar.append(SatirHatasi(satir_no, _kimlik_alani(tur),
                                             f"kimlik dosyada tekrar ediyor (ilk: satır {gorulen[kayit.dis_kimlik]})"))
            continue
        gorulen[kayit.dis_kimlik] = satir_no
        h.kayitlar.append(kayit)
    if tur == "uyeler":
        h.kayitlar = _telefon_tekillestir(h.kayitlar, h.uyarilar)
    elif tur == "paketler":
        h.kayitlar = _zincirle(h.kayitlar)
    return h


@dataclass(frozen=True)
class _Baglam:
    dilim: ZoneInfo
    bugun: date
    simdi: datetime
    anonim: bool
    dis_kaynak: str
    bilinen_uyeler: set[str] | None


def _kimlik_alani(tur: str) -> str:
    return {"uyeler": "uye_kimlik", "paketler": "paket_kimlik", "girisler": "giris_kimlik"}[tur]


def _alan(alan: str, hatalar: list[SatirHatasi], satir_no: int, fonk, *argumanlar):
    """fonk(*argumanlar); DegerHatasi'nı satır hatasına çevirir (değer olmadan) ve None döndürür."""
    try:
        return fonk(*argumanlar)
    except DegerHatasi as hata:
        hatalar.append(SatirHatasi(satir_no, alan, str(hata)))
        return None


def _zorunlu_metin(deger) -> str:
    metin = _metin(deger)
    if metin is None:
        raise DegerHatasi("boş")
    return metin


def _uye_kimligi(satir_no, deger, b: _Baglam, hatalar) -> str | None:
    kimlik = _alan("uye_kimlik", hatalar, satir_no, _zorunlu_metin, deger.get("uye_kimlik"))
    if kimlik is not None and b.bilinen_uyeler is not None and kimlik not in b.bilinen_uyeler:
        hatalar.append(SatirHatasi(satir_no, "uye_kimlik", "üye bulunamadı (ne üye dosyasında ne veritabanında)"))
    return kimlik


def _uye(satir_no, deger, b: _Baglam, hatalar, uyarilar) -> UyeKaydi | None:
    kimlik = _alan("uye_kimlik", hatalar, satir_no, _zorunlu_metin, deger.get("uye_kimlik"))
    if kimlik is None:
        return None
    if b.anonim:                                 # K35: ad, telefon ve e-posta hiç okunmaz
        return UyeKaydi(satir_no, kimlik, anonim_ad(b.dis_kaynak, kimlik), None, None)
    ad_soyad = _metin(deger.get("ad_soyad"))
    if ad_soyad is None:
        ad_soyad = " ".join(p for p in (_metin(deger.get("ad")), _metin(deger.get("soyad"))) if p) or None
    if ad_soyad is None:
        hatalar.append(SatirHatasi(satir_no, "ad_soyad", "boş"))
        return None
    telefon = None
    if (ham := _metin(deger.get("telefon"))) is not None:
        try:
            telefon = telefon_normallestir(ham)
        except ValueError:
            uyarilar["telefon_gecersiz"] += 1
    eposta = _metin(deger.get("eposta"))
    if eposta is not None and "@" not in eposta:
        uyarilar["eposta_gecersiz"] += 1
        eposta = None
    return UyeKaydi(satir_no, kimlik, ad_soyad, telefon, eposta)


def _paket(satir_no, deger, b: _Baglam, hatalar, uyarilar) -> PaketKaydi | None:
    uye = _uye_kimligi(satir_no, deger, b, hatalar)
    ad = _alan("paket_adi", hatalar, satir_no, _zorunlu_metin, deger.get("paket_adi"))
    baslangic = _alan("baslangic", hatalar, satir_no, tarih_coz, deger.get("baslangic"))
    bitis = None
    if deger.get("bitis") is not None:
        bitis = _alan("bitis", hatalar, satir_no, tarih_coz, deger.get("bitis"))
    hak = None
    if deger.get("giris_hakki") is not None:
        hak = _alan("giris_hakki", hatalar, satir_no, _tam_sayi, deger.get("giris_hakki"))
        if hak is not None and hak <= 0:
            hatalar.append(SatirHatasi(satir_no, "giris_hakki", "sıfırdan büyük olmalı"))
    ucret = _alan("ucret", hatalar, satir_no, para_coz, deger.get("ucret"), uyarilar)
    durum = _alan("durum", hatalar, satir_no, durum_coz, "paketler", deger.get("durum"))
    if hatalar:
        return None
    baslangic, bitis = baslangic[0], (bitis[0] if bitis else None)
    tur = "giris" if hak is not None else "sure"
    if tur == "sure" and bitis is None:
        hatalar.append(SatirHatasi(satir_no, "bitis", "süre bazlı pakette (giriş hakkı yok) zorunlu"))
    if bitis is not None and bitis <= baslangic:
        hatalar.append(SatirHatasi(satir_no, "bitis", "başlangıçtan sonra olmalı"))
    if hatalar:
        return None
    if durum is None:
        durum = "aktif" if bitis is None or bitis >= b.bugun else "bitti"
    kimlik = _metin(deger.get("paket_kimlik")) or "P" + _ozet(uye, baslangic.isoformat(), ad)[:16]
    return PaketKaydi(satir_no, kimlik, uye, tur, ad, baslangic, bitis, hak, ucret, durum, None)


def _giris(satir_no, deger, b: _Baglam, hatalar, uyarilar) -> GirisKaydi | None:
    uye = _uye_kimligi(satir_no, deger, b, hatalar)
    gun_saat = _alan("tarih", hatalar, satir_no, tarih_coz, deger.get("tarih"))
    saat = None
    if deger.get("saat") is not None:
        saat = _alan("saat", hatalar, satir_no, saat_coz, deger.get("saat"))
    zaman = None
    if not hatalar:
        zaman = _alan("tarih", hatalar, satir_no, _yerel_zaman, gun_saat, saat, b.dilim, b.simdi)
    durum = _alan("durum", hatalar, satir_no, durum_coz, "girisler", deger.get("durum"))
    tutar = Decimal("0")
    if deger.get("tutar") is not None:
        tutar = _alan("tutar", hatalar, satir_no, para_coz, deger.get("tutar"), uyarilar)
    if hatalar:
        return None
    yerel = zaman.astimezone(b.dilim).replace(tzinfo=None).isoformat()
    kimlik = _metin(deger.get("giris_kimlik")) or "G" + _ozet(uye, yerel)[:16]
    return GirisKaydi(satir_no, kimlik, uye, zaman, durum, tutar)


def _telefon_tekillestir(kayitlar: list[UyeKaydi], uyarilar: Counter) -> list[UyeKaydi]:
    """Veritabanında (isletme_id, telefon_e164) tekildir: aynı telefon ikinci üyede boş bırakılır."""
    gorulen: set[str] = set()
    sonuc = []
    for k in kayitlar:
        if k.telefon_e164 is not None:
            if k.telefon_e164 in gorulen:
                uyarilar["telefon_tekrar"] += 1
                k = replace(k, telefon_e164=None)
            else:
                gorulen.add(k.telefon_e164)
        sonuc.append(k)
    return sonuc


def _zincirle(kayitlar: list[PaketKaydi]) -> list[PaketKaydi]:
    """Her üyenin paketleri başlangıca (eşitlikte satır sırasına) göre sıralanır; her paket bir öncekine bağlanır."""
    onceki: dict[str, str] = {}
    gruplar: dict[str, list[PaketKaydi]] = defaultdict(list)
    for k in kayitlar:
        gruplar[k.uye_kimlik].append(k)
    for liste in gruplar.values():
        liste.sort(key=lambda k: (k.baslangic_tarihi, k.satir))
        for once, sonra in zip(liste, liste[1:]):
            onceki[sonra.dis_kimlik] = once.dis_kimlik
    return [replace(k, onceki_dis_kimlik=onceki.get(k.dis_kimlik)) for k in kayitlar]
